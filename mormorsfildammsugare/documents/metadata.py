"""Document metadata & text extractors for PDF, Office, Emails, and Text.

Ported and enhanced from Arkivslukaren.
"""

from dataclasses import dataclass, field
from datetime import datetime, timezone
import email
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import List, Optional
import re


@dataclass
class DocumentMetadata:
    created_date: Optional[datetime] = None
    title: Optional[str] = None
    subject: Optional[str] = None
    author: Optional[str] = None
    keywords: List[str] = field(default_factory=list)
    sheet_names: List[str] = field(default_factory=list)
    text_sample: Optional[str] = None


def parse_pdf(filepath: Path) -> DocumentMetadata:
    """Extract metadata and first-page text from PDF using pypdf."""
    meta = DocumentMetadata()
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(filepath))
        info = reader.metadata
        if info:
            meta.title = info.title
            meta.subject = info.subject
            meta.author = info.author
            if info.get("/Keywords"):
                meta.keywords = [k.strip() for k in str(info.get("/Keywords")).split(",") if k.strip()]

            # CreationDate format: D:YYYYMMDDHHmmSS...
            c_date = info.get("/CreationDate")
            if c_date and isinstance(c_date, str) and c_date.startswith("D:"):
                try:
                    s = c_date[2:16]
                    if len(s) >= 8:
                        y = int(s[0:4])
                        m = int(s[4:6])
                        d = int(s[6:8])
                        meta.created_date = datetime(y, m, d)
                except Exception:
                    pass

        # Text sample (first page, up to 4KB)
        if len(reader.pages) > 0:
            txt = reader.pages[0].extract_text()
            if txt:
                meta.text_sample = txt[:4000]
    except Exception:
        pass
    return meta


def parse_docx(filepath: Path) -> DocumentMetadata:
    """Extract metadata and text from DOCX using python-docx."""
    meta = DocumentMetadata()
    try:
        import docx
        doc = docx.Document(str(filepath))
        cp = doc.core_properties
        meta.title = cp.title
        meta.subject = cp.subject
        meta.author = cp.author
        if cp.keywords:
            meta.keywords = [k.strip() for k in cp.keywords.split(",") if k.strip()]
        if cp.created:
            meta.created_date = cp.created.replace(tzinfo=None)

        # Paragraph text sample
        lines = []
        for p in doc.paragraphs[:20]:
            if p.text:
                lines.append(p.text)
                if sum(len(l) for l in lines) > 4000:
                    break
        if lines:
            meta.text_sample = "\n".join(lines)[:4000]
    except Exception:
        pass
    return meta


def parse_xlsx(filepath: Path) -> DocumentMetadata:
    """Extract metadata and sheet names from XLSX using openpyxl."""
    meta = DocumentMetadata()
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(filepath), read_only=True)
        meta.sheet_names = wb.sheetnames
        props = wb.properties
        if props:
            meta.title = props.title
            meta.subject = props.subject
            meta.author = props.creator
            if props.keywords:
                meta.keywords = [k.strip() for k in props.keywords.split(",") if k.strip()]
            if props.created:
                meta.created_date = props.created.replace(tzinfo=None)
        wb.close()
    except Exception:
        pass
    return meta


def parse_eml(filepath: Path) -> DocumentMetadata:
    """Extract Date, Subject, and From from RFC 822 email message."""
    meta = DocumentMetadata()
    try:
        with open(filepath, "rb") as f:
            msg = email.message_from_binary_file(f)

        meta.title = msg.get("Subject")
        meta.author = msg.get("From")

        date_str = msg.get("Date")
        if date_str:
            try:
                dt = parsedate_to_datetime(date_str)
                meta.created_date = dt.astimezone(timezone.utc).replace(tzinfo=None)
            except Exception:
                pass

        # Text snippet from payload
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    try:
                        meta.text_sample = part.get_payload(decode=True).decode("utf-8", errors="ignore")[:4000]
                        break
                    except Exception:
                        pass
        else:
            try:
                meta.text_sample = msg.get_payload(decode=True).decode("utf-8", errors="ignore")[:4000]
            except Exception:
                pass
    except Exception:
        pass
    return meta


def parse_text(filepath: Path) -> DocumentMetadata:
    """Read first 4KB of plain text / markdown / CSV."""
    meta = DocumentMetadata()
    try:
        with open(filepath, "rb") as f:
            raw = f.read(4096)
        # Try UTF-8 first, then Latin-1
        try:
            meta.text_sample = raw.decode("utf-8")
        except UnicodeDecodeError:
            meta.text_sample = raw.decode("latin-1", errors="ignore")
    except Exception:
        pass
    return meta


def extract_document_metadata(filepath: Path) -> DocumentMetadata:
    """Unified entry point for document metadata extraction."""
    ext = filepath.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(filepath)
    if ext in (".docx", ".dotx", ".docm"):
        return parse_docx(filepath)
    if ext in (".xlsx", ".xlsm", ".xltx"):
        return parse_xlsx(filepath)
    if ext == ".eml":
        return parse_eml(filepath)
    if ext in (".txt", ".csv", ".tsv", ".md", ".rtf"):
        return parse_text(filepath)
    return DocumentMetadata()
